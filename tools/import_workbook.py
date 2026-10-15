"""Bring a completed ISV Discovery Workbook into the pipeline.

    python tools/import_workbook.py Customer_ISV_Discovery_Workbook.xlsx [--project samples/01-isv-inventory]

Reads only what the customer typed (the yellow cells) — never the calculated columns, so it works on a
file saved by Excel, Google Sheets or Numbers without recalculation. Columns are found by header text,
so reordering or hiding columns is fine. Then it does what tools/import_intake.py does: writes
data/intake_<customer>.csv and adds the customer and cloud commitments to reference/customers.yaml.
The example row (notes start with "EXAMPLE ROW") is skipped.
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from import_intake import write_intake  # noqa: E402

# workbook header (without " *") -> intake line key
LINE_MAP = {
    "Product / application": "product", "Edition / SKU / plan": "edition", "Vendor": "vendor",
    "Department / business unit": "department", "Business owner": "owner", "Primary use case": "useCase",
    "Business criticality": "criticality", "Licenses / seats": "licenses", "Active users (last 90 days)": "activeUsers",
    "Price per user / month ($)": "unitMonthly", "Annual cost ($)": "annualCost", "Purchased via": "purchasedVia",
    "Renewal / contract end date": "renewalDate", "Term (months)": "termMonths", "Auto-renews?": "autoRenew",
    "Notice period (days)": "noticeDays", "Satisfaction (1–5)": "satisfaction",
    "Open to alternatives after renewal?": "openToAlternatives", "Alternatives being considered": "evaluating",
    "Main reason to switch": "switchReason", "Desired outcome at renewal": "desiredOutcome", "Notes": "notes",
    "Pain points / gaps": "_pain", "Other use cases": "_other",
}
PROFILE_MAP = {
    "Company legal name": ("customer", "name"), "Other names used (comma-separated)": ("customer", "aliases"),
    "Industry": ("customer", "industry"), "Employees": ("customer", "employees"),
    "Executive sponsor": ("customer", "sponsorName"), "Sponsor title": ("customer", "sponsorTitle"),
    "Consultant / account lead": ("opportunity", "accountLead"),
    "AWS commitment (EDP / private pricing)?": ("aws", "has"), "AWS annual commitment ($)": ("aws", "annual"),
    "AWS commitment remaining this period ($)": ("aws", "remaining"), "AWS commitment period ends": ("aws", "end"),
    "Microsoft commitment (MACC)?": ("azure", "has"), "Microsoft annual commitment ($)": ("azure", "annual"),
    "Microsoft commitment remaining this period ($)": ("azure", "remaining"), "Microsoft commitment period ends": ("azure", "end"),
}


def clean(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def money(v):
    if v is None or isinstance(v, (int, float)):
        return v
    s = re.sub(r"[,$\s]", "", str(v)).lower()
    m = re.match(r"^(\d+(?:\.\d+)?)([km]?)$", s)
    return float(m.group(1)) * {"": 1, "k": 1e3, "m": 1e6}[m.group(2)] if m else None


def read_workbook(path: Path) -> dict:
    wb = load_workbook(path, data_only=True)
    if "ISV Inventory" not in wb.sheetnames or "Company Profile" not in wb.sheetnames:
        raise SystemExit(f"{path} doesn't look like an ISV Discovery Workbook (needs 'ISV Inventory' and 'Company Profile' tabs)")
    intake = {"opportunity": {"source": "Discovery workbook"}, "customer": {}, "commitments": {"aws": {}, "azure": {}}, "lines": []}

    pr = wb["Company Profile"]
    for row in pr.iter_rows(min_row=1, max_col=3):
        label = (row[1].value or "").replace("*", "").strip() if len(row) > 1 else ""
        if label in PROFILE_MAP:
            grp, key = PROFILE_MAP[label]
            val = clean(row[2].value)
            if key in ("annual", "remaining"):
                val = money(val)
            target = intake["commitments"][grp] if grp in ("aws", "azure") else intake[grp]
            if val is not None:
                target[key] = val

    inv = wb["ISV Inventory"]
    header_row = next(r for r in range(1, 15) if any(str(c.value or "").startswith("Product / application") for c in inv[r]))
    cols = {}
    for c in inv[header_row]:
        h = str(c.value or "").replace(" *", "").strip()
        if h in LINE_MAP:
            cols[LINE_MAP[h]] = c.column - 1
    skipped = 0
    for row in inv.iter_rows(min_row=header_row + 1, values_only=True):
        line = {k: clean(row[i]) for k, i in cols.items() if i < len(row)}
        if not line.get("product"):
            continue
        if str(line.get("notes") or "").upper().startswith("EXAMPLE ROW"):
            skipped += 1
            continue
        for k in ("annualCost", "unitMonthly"):
            line[k] = money(line.get(k))
        if line.get("satisfaction") is not None:
            line["satisfaction"] = str(int(float(line["satisfaction"])))
        pain, other = line.pop("_pain", None), line.pop("_other", None)
        notes = [line.get("notes")] + ([f"Other uses: {other}"] if other else []) + ([f"Pain points: {pain}"] if pain else [])
        line["notes"] = " · ".join(n for n in notes if n) or None
        intake["lines"].append(line)
    if skipped:
        print(f"skipped {skipped} example row(s)")
    return intake


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--project", type=Path, default=Path("samples/01-isv-inventory"))
    a = ap.parse_args(argv)
    return write_intake(read_workbook(a.workbook), a.project)


if __name__ == "__main__":
    sys.exit(main())
