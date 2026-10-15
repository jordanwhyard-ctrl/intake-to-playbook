"""Build the ISV Discovery Workbook: the spreadsheet a consultant sends a customer (or fills in with them)
to document every ISV product they pay for, its use cases, renewal terms, and openness to alternatives.

    python tools/build_discovery_workbook.py                 # -> templates/ISV_Discovery_Workbook.xlsx
    python tools/build_discovery_workbook.py --example       # -> templates/ISV_Discovery_Workbook_EXAMPLE.xlsx

The product dropdown, catalog lookups and marketplace columns come from
samples/01-isv-inventory/reference/isv_catalog.yaml, so rebuild after changing the catalog.
A completed workbook goes into the pipeline with tools/import_workbook.py.
"""
from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "samples/01-isv-inventory/reference/isv_catalog.yaml"

# ---------------------------------------------------------------- style
FONT = "Arial"
INK = "1F2A27"
TEAL = "0D6B5C"
TEAL_SOFT = "D5EBE5"
INPUT = "FFF7D6"      # cells the customer fills in
CALC = "EEF1F0"       # calculated, don't edit
BAND = {"Product": "0D6B5C", "Ownership & use case": "2F5D7C", "Licenses & cost": "6B4E9B",
        "Contract & renewal": "A35A00", "Satisfaction & alternatives": "9B2C48", "From the catalog (automatic)": "5A6A65"}
thin = Side(style="thin", color="D4DBD8")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def f(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


# ---------------------------------------------------------------- dropdown lists
LISTS = {
    "YesNo": ["Yes", "No", "Unknown"],
    "Criticality": ["Mission-critical", "High", "Medium", "Low"],
    "Sensitivity": ["Public", "Internal", "Confidential", "Regulated (PII / PHI / PCI)"],
    "Billing": ["Annual upfront", "Monthly", "Quarterly", "Multi-year upfront", "Usage-based", "Unknown"],
    "Channel": ["Direct", "Reseller", "AWS Marketplace", "Microsoft Marketplace", "Google Cloud Marketplace", "Unknown"],
    "Satisfaction": ["1", "2", "3", "4", "5"],
    "Openness": ["Yes – actively evaluating", "Yes – open at renewal", "Maybe – depends on price", "No – committed", "No – contractually locked in", "Unknown"],
    "Reason": ["Cost", "Missing features", "Poor support", "Overlaps another tool", "Security / compliance", "Low adoption", "Vendor viability", "Consolidation mandate", "No reason to switch", "Other"],
    "Outcome": ["Renew as-is", "Renegotiate price", "Right-size licenses", "Upgrade / expand", "Consolidate onto another tool", "Replace with alternative", "Retire", "Undecided"],
    "Industry": ["Financial services", "Healthcare", "Legal", "Logistics", "Manufacturing", "Retail", "Software", "Education", "Public sector", "Energy", "Media", "Professional services", "Other"],
    "Months": ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"],
    "Counts": ["Yes", "Partly", "No", "Unknown"],
}

# ---------------------------------------------------------------- inventory columns
# (key, header, section, width, kind, list/format, required, help)
COLS = [
    ("n", "#", "Product", 5, "calc", None, False, None),
    ("product", "Product / application", "Product", 28, "in", "product", True, "Pick from the list or type the name as it appears on the invoice."),
    ("edition", "Edition / SKU / plan", "Product", 18, "in", None, False, "e.g. Enterprise, E3, Business+"),
    ("vendor", "Vendor", "Product", 16, "in", None, False, "Only needed if the product isn't in the list."),
    ("department", "Department / business unit", "Ownership & use case", 18, "in", None, False, None),
    ("owner", "Business owner", "Ownership & use case", 18, "in", None, False, "Person or team accountable for the tool."),
    ("tech_owner", "Technical owner / admin", "Ownership & use case", 18, "in", None, False, None),
    ("use_case", "Primary use case", "Ownership & use case", 32, "in", None, True, "What the business uses it for, in a sentence."),
    ("other_uses", "Other use cases", "Ownership & use case", 26, "in", None, False, None),
    ("criticality", "Business criticality", "Ownership & use case", 15, "in", "Criticality", False, None),
    ("integrations", "Key integrations", "Ownership & use case", 24, "in", None, False, "Systems it connects to (SSO, CRM, ERP, data warehouse…)"),
    ("sensitivity", "Data sensitivity", "Ownership & use case", 18, "in", "Sensitivity", False, None),
    ("licenses", "Licenses / seats", "Licenses & cost", 11, "in", "int", False, None),
    ("active", "Active users (last 90 days)", "Licenses & cost", 13, "in", "int", False, "From the admin console, if available."),
    ("util", "Utilization", "Licenses & cost", 11, "calc", "pct", False, None),
    ("unit", "Price per user / month ($)", "Licenses & cost", 13, "in", "money2", False, "Give this + licenses, or the annual cost."),
    ("annual_cost", "Annual cost ($)", "Licenses & cost", 14, "in", "money", False, "Total paid per year. Leave blank if you gave price per user."),
    ("annual_spend", "Annual spend ($)", "Licenses & cost", 14, "calc", "money", False, None),
    ("billing", "Billing frequency", "Licenses & cost", 15, "in", "Billing", False, None),
    ("channel", "Purchased via", "Licenses & cost", 18, "in", "Channel", False, None),
    ("start", "Contract start date", "Contract & renewal", 13, "in", "date", False, None),
    ("renewal", "Renewal / contract end date", "Contract & renewal", 14, "in", "date", True, None),
    ("term", "Term (months)", "Contract & renewal", 10, "in", "int", False, None),
    ("auto", "Auto-renews?", "Contract & renewal", 11, "in", "YesNo", False, None),
    ("notice", "Notice period (days)", "Contract & renewal", 11, "in", "int", False, "Days before renewal to give notice of non-renewal."),
    ("notice_by", "Notice deadline", "Contract & renewal", 13, "calc", "date", False, None),
    ("days", "Days to renewal", "Contract & renewal", 10, "calc", "int", False, None),
    ("window", "Renewal window", "Contract & renewal", 14, "calc", None, False, None),
    ("cap", "Price increase cap at renewal (%)", "Contract & renewal", 13, "in", "pct", False, "Contractual cap on uplift, if any."),
    ("tfc", "Can terminate early?", "Contract & renewal", 12, "in", "YesNo", False, "Termination for convenience."),
    ("sat", "Satisfaction (1–5)", "Satisfaction & alternatives", 11, "in", "Satisfaction", False, "5 = very satisfied"),
    ("pain", "Pain points / gaps", "Satisfaction & alternatives", 28, "in", None, False, None),
    ("open", "Open to alternatives after renewal?", "Satisfaction & alternatives", 22, "in", "Openness", True, None),
    ("considering", "Alternatives being considered", "Satisfaction & alternatives", 24, "in", None, False, None),
    ("reason", "Main reason to switch", "Satisfaction & alternatives", 18, "in", "Reason", False, None),
    ("barriers", "Switching barriers", "Satisfaction & alternatives", 24, "in", None, False, "Data migration, retraining, integrations, contract…"),
    ("outcome", "Desired outcome at renewal", "Satisfaction & alternatives", 20, "in", "Outcome", False, None),
    ("decider", "Decision maker", "Satisfaction & alternatives", 18, "in", None, False, None),
    ("notes", "Notes", "Satisfaction & alternatives", 28, "in", None, False, None),
    ("matched", "Catalog product", "From the catalog (automatic)", 24, "calc", None, False, None),
    ("in_cat", "In catalog?", "From the catalog (automatic)", 10, "calc", None, False, None),
    ("cat_vendor", "Catalog vendor", "From the catalog (automatic)", 16, "calc", None, False, None),
    ("cat_category", "Category", "From the catalog (automatic)", 18, "calc", None, False, None),
    ("aws", "AWS Marketplace", "From the catalog (automatic)", 13, "calc", None, False, None),
    ("ms", "Microsoft Marketplace", "From the catalog (automatic)", 13, "calc", None, False, None),
    ("alts", "Catalog alternatives", "From the catalog (automatic)", 34, "calc", None, False, None),
    ("key", "sort key", "From the catalog (automatic)", 4, "calc", None, False, None),
]
C = {k: get_column_letter(i + 1) for i, (k, *_r) in enumerate(COLS)}
HEAD_ROW, FIRST, NROWS = 5, 6, 250
LAST = FIRST + NROWS - 1
STATUS = {"listed": "Listed", "app": "Add-in only", "not_found": "Not found", "n_a": "n/a", "first_party": "Microsoft", "unverified": "Unverified"}


def formula(k: str, r: int) -> str | None:
    p, c = f"{C['product']}{r}", C
    m = f"{C['matched']}{r}"
    cat = "Catalog!$A$5:$A$400"
    look = lambda col: f'=IF({m}="","",IFERROR(INDEX(Catalog!${col}$5:${col}$400,MATCH({m},{cat},0)),""))'
    return {
        "matched": f'=IF({p}="","",IFERROR(INDEX(Lookup!$B$2:$B$3000,MATCH(TRIM({p}),Lookup!$A$2:$A$3000,0)),""))',
        "n": f'=IF({p}="","",ROW()-{FIRST - 1})',
        "util": f'=IF(OR({c["licenses"]}{r}="",{c["active"]}{r}="",{c["licenses"]}{r}=0),"",{c["active"]}{r}/{c["licenses"]}{r})',
        "annual_spend": f'=IF({c["annual_cost"]}{r}<>"",{c["annual_cost"]}{r},IF(AND({c["unit"]}{r}<>"",{c["licenses"]}{r}<>""),{c["unit"]}{r}*12*{c["licenses"]}{r},""))',
        "notice_by": f'=IF(OR({c["renewal"]}{r}="",{c["notice"]}{r}=""),"",{c["renewal"]}{r}-{c["notice"]}{r})',
        "days": f'=IF({c["renewal"]}{r}="","",{c["renewal"]}{r}-TODAY())',
        "window": (f'=IF({c["renewal"]}{r}="","",IF({c["days"]}{r}<0,"Past due",IF({c["days"]}{r}<=90,"0–90 days",'
                   f'IF({c["days"]}{r}<=180,"91–180 days",IF({c["days"]}{r}<=365,"181–365 days","Over a year")))))'),
        "in_cat": f'=IF({p}="","",IF({m}="","No – will be reviewed","Yes"))',
        "cat_vendor": look("B"),
        "cat_category": look("C"),
        "aws": look("D"),
        "ms": look("E"),
        "alts": look("F"),
        "key": f'=IF(AND({p}<>"",ISNUMBER({c["renewal"]}{r})),{c["renewal"]}{r}+ROW()/100000,"")',
    }.get(k)


NUMFMT = {"money": '$#,##0;($#,##0);"-"', "money2": '$#,##0.00', "pct": "0%", "int": "#,##0", "date": "yyyy-mm-dd"}

EXAMPLE_ROW = {
    "product": "Okta Workforce Identity", "edition": "SSO + MFA", "department": "IT", "owner": "IT Director",
    "tech_owner": "Identity team", "use_case": "Single sign-on and MFA for all employees and contractors",
    "other_uses": "Lifecycle provisioning to HR system", "criticality": "Mission-critical",
    "integrations": "Workday, Google Workspace, AWS", "sensitivity": "Confidential", "licenses": 650, "active": 612,
    "unit": None, "annual_cost": 58500, "billing": "Annual upfront", "channel": "Direct", "start": date(2025, 11, 21),
    "renewal": date(2026, 11, 20), "term": 12, "auto": "Yes", "notice": 30, "cap": 0.05, "tfc": "No", "sat": "4",
    "pain": "Price went up 7% last year", "open": "Maybe – depends on price", "considering": "Microsoft Entra ID P1",
    "reason": "Cost", "barriers": "120 app integrations to re-point", "outcome": "Renegotiate price", "decider": "CFO",
    "notes": "EXAMPLE ROW: replace or delete",
}

FULL_EXAMPLE = [  # fictional customer, used only by --example
    dict(EXAMPLE_ROW, notes=""),
    {"product": "Datadog", "edition": "Enterprise", "department": "Engineering", "owner": "VP Engineering", "use_case": "Infrastructure and APM monitoring for production services", "criticality": "Mission-critical", "integrations": "AWS, PagerDuty, Slack", "sensitivity": "Internal", "annual_cost": 420000, "billing": "Annual upfront", "channel": "Direct", "renewal": date(2026, 12, 31), "term": 12, "auto": "No", "sat": "4", "pain": "Host-based pricing grows faster than usage", "open": "Yes – open at renewal", "considering": "Grafana Cloud, New Relic", "reason": "Cost", "barriers": "Dashboards and monitors to rebuild", "outcome": "Renegotiate price", "decider": "VP Engineering"},
    {"product": "Splunk Enterprise Security", "edition": "Cloud", "department": "Security", "owner": "CISO", "use_case": "SIEM and security log search", "criticality": "Mission-critical", "integrations": "CrowdStrike, Okta, AWS CloudTrail", "sensitivity": "Regulated (PII / PHI / PCI)", "annual_cost": 310000, "billing": "Annual upfront", "channel": "Direct", "renewal": date(2027, 6, 30), "term": 36, "auto": "No", "sat": "2", "pain": "Ingest costs keep climbing", "open": "Yes – actively evaluating", "considering": "Sumo Logic, Elastic Cloud", "reason": "Cost", "barriers": "Detection rules and dashboards", "outcome": "Replace with alternative", "decider": "CISO"},
    {"product": "Notion", "edition": "Business", "department": "Product", "owner": "VP Product", "use_case": "Product specs and team wikis", "criticality": "Medium", "sensitivity": "Internal", "licenses": 600, "active": 210, "unit": 10, "billing": "Annual upfront", "channel": "Direct", "renewal": date(2027, 1, 31), "term": 12, "auto": "Yes", "notice": 30, "sat": "3", "pain": "Most docs already live in Confluence", "open": "Yes – open at renewal", "considering": "Atlassian Confluence", "reason": "Overlaps another tool", "barriers": "Page migration", "outcome": "Consolidate onto another tool", "decider": "CTO"},
    {"product": "Atlassian Confluence", "edition": "Premium", "department": "All", "owner": "IT Director", "use_case": "Company knowledge base and engineering docs", "criticality": "High", "integrations": "Jira", "sensitivity": "Internal", "licenses": 600, "active": 590, "annual_cost": 72000, "billing": "Annual upfront", "channel": "Direct", "renewal": date(2027, 3, 31), "term": 12, "auto": "No", "sat": "4", "open": "No – committed", "outcome": "Upgrade / expand", "decider": "CTO"},
    {"product": "GitLab Ultimate", "department": "Engineering", "owner": "VP Engineering", "use_case": "Source control, CI/CD and security scanning", "criticality": "Mission-critical", "integrations": "AWS, Jira, Snyk", "sensitivity": "Confidential", "licenses": 220, "active": 140, "annual_cost": 217800, "billing": "Annual upfront", "channel": "Direct", "renewal": date(2026, 12, 15), "term": 12, "auto": "Yes", "notice": 45, "sat": "3", "pain": "Contractors offboarded; seats unused", "open": "Maybe – depends on price", "considering": "GitHub Enterprise", "reason": "Low adoption", "barriers": "Pipelines and runners", "outcome": "Right-size licenses", "decider": "VP Engineering"},
    {"product": "Snyk", "edition": "Enterprise", "department": "Engineering", "owner": "Platform lead", "use_case": "Dependency and container vulnerability scanning", "criticality": "High", "integrations": "GitLab", "sensitivity": "Internal", "licenses": 200, "active": 190, "annual_cost": 88000, "billing": "Annual upfront", "channel": "Direct", "renewal": date(2027, 2, 28), "term": 12, "auto": "No", "sat": "2", "pain": "False positives; developers ignore findings", "open": "Yes – actively evaluating", "considering": "Checkmarx One", "reason": "Missing features", "outcome": "Replace with alternative", "decider": "CISO"},
    {"product": "Zoom", "edition": "Business", "department": "All", "owner": "IT Director", "use_case": "Video meetings with customers", "criticality": "Medium", "sensitivity": "Internal", "licenses": 600, "active": 480, "annual_cost": 79200, "billing": "Annual upfront", "channel": "Direct", "renewal": date(2026, 12, 31), "term": 12, "auto": "Yes", "notice": 30, "sat": "3", "pain": "Most meetings moved to Google Meet", "open": "Yes – open at renewal", "considering": "Google Meet", "reason": "Overlaps another tool", "outcome": "Retire", "decider": "IT Director"},
    {"product": "Snowflake", "edition": "Enterprise", "department": "Data", "owner": "Head of Data", "use_case": "Data warehouse for analytics and finance reporting", "criticality": "Mission-critical", "integrations": "Fivetran, dbt, Tableau", "sensitivity": "Confidential", "annual_cost": 610000, "billing": "Usage-based", "channel": "AWS Marketplace", "renewal": date(2027, 1, 31), "term": 12, "auto": "No", "sat": "4", "open": "No – committed", "outcome": "Renegotiate price", "decider": "CFO"},
    {"product": "MatterDesk Pro", "department": "Legal", "owner": "General Counsel", "use_case": "Contract intake and matter tracking", "criticality": "Medium", "sensitivity": "Confidential", "licenses": 15, "active": 14, "annual_cost": 18000, "billing": "Annual upfront", "channel": "Direct", "renewal": date(2027, 9, 30), "term": 12, "auto": "Yes", "notice": 60, "sat": "4", "open": "No – committed", "outcome": "Renew as-is", "decider": "General Counsel", "notes": "Niche tool, not in the catalog"},
]

PROFILE = [
    ("Company", None),
    ("Company legal name *", "name"), ("Other names used (comma-separated)", "aliases"), ("Industry *", "industry"),
    ("Employees *", "employees"), ("Headquarters country", "country"), ("Fiscal year ends (month)", "fy_end"),
    ("People", None),
    ("Executive sponsor *", "sponsor"), ("Sponsor title", "sponsor_title"), ("Day-to-day contact", "contact"),
    ("Contact email", "contact_email"), ("Procurement contact", "procurement"), ("Security / compliance contact", "security"),
    ("How you buy software", None),
    ("Buy direct from vendors?", "buy_direct"), ("Buy through resellers?", "buy_reseller"), ("Buy on AWS Marketplace?", "buy_aws"),
    ("Buy on Microsoft Marketplace?", "buy_ms"), ("Private marketplace enabled?", "private_mp"),
    ("New vendors need a security review?", "sec_review"), ("Typical procurement lead time (days)", "proc_days"),
    ("Approval needed above ($)", "approval_threshold"),
    ("Cloud commitments", None),
    ("AWS commitment (EDP / private pricing)?", "aws_has"), ("AWS annual commitment ($)", "aws_annual"),
    ("AWS commitment remaining this period ($)", "aws_remaining"), ("AWS commitment period ends", "aws_end"),
    ("Do AWS Marketplace purchases count toward it?", "aws_counts"),
    ("Microsoft commitment (MACC)?", "ms_has"), ("Microsoft annual commitment ($)", "ms_annual"),
    ("Microsoft commitment remaining this period ($)", "ms_remaining"), ("Microsoft commitment period ends", "ms_end"),
    ("Do Microsoft Marketplace purchases count toward it?", "ms_counts"),
    ("This workbook", None),
    ("Completed by", "completed_by"), ("Date completed", "completed_on"), ("Consultant / account lead", "account_lead"),
]
PROFILE_LISTS = {"industry": "Industry", "fy_end": "Months", "buy_direct": "YesNo", "buy_reseller": "YesNo", "buy_aws": "YesNo",
                 "buy_ms": "YesNo", "private_mp": "YesNo", "sec_review": "YesNo", "aws_has": "YesNo", "aws_counts": "Counts",
                 "ms_has": "YesNo", "ms_counts": "Counts"}
PROFILE_FMT = {"employees": "int", "proc_days": "int", "approval_threshold": "money", "aws_annual": "money",
               "aws_remaining": "money", "aws_end": "date", "ms_annual": "money", "ms_remaining": "money", "ms_end": "date",
               "completed_on": "date"}
PROFILE_EXAMPLE = {"name": "Example Software Co. (fictional)", "aliases": "Example Software, ESC", "industry": "Software",
                   "employees": 600, "country": "United States", "fy_end": "December", "sponsor": "Jamie Rivera",
                   "sponsor_title": "CFO", "contact": "Sam Patel", "contact_email": "sam.patel@example.com",
                   "buy_direct": "Yes", "buy_reseller": "No", "buy_aws": "Yes", "buy_ms": "No", "private_mp": "No",
                   "sec_review": "Yes", "proc_days": 30, "approval_threshold": 50000, "aws_has": "Yes", "aws_annual": 2500000,
                   "aws_remaining": 1400000, "aws_end": date(2027, 1, 31), "aws_counts": "Yes", "ms_has": "No",
                   "account_lead": "Marcus Lee"}

QUESTIONS = [
    ("Strategy", "What business goals does software spend need to support over the next 12–24 months?"),
    ("Strategy", "Is there a mandate to consolidate vendors or cut software spend? By how much?"),
    ("Strategy", "Are any reorganizations, acquisitions or divestitures planned that change who uses what?"),
    ("Renewals", "Who owns the renewal calendar today, and how much notice do you usually get before a renewal?"),
    ("Renewals", "Which renewals in the next 6 months worry you most, and why?"),
    ("Renewals", "Have any auto-renewals caught you by surprise in the last year?"),
    ("Renewals", "What price increases have vendors proposed at recent renewals?"),
    ("Use cases", "Where do two or more tools do the same job? Which would you keep?"),
    ("Use cases", "Which tools were bought by one team and spread without central approval?"),
    ("Use cases", "Which capabilities are missing today that you're working around?"),
    ("Alternatives", "Which products would you replace tomorrow if switching cost nothing?"),
    ("Alternatives", "What usually stops you from switching: data migration, retraining, integrations, or contract terms?"),
    ("Alternatives", "Are any vendors on a preferred or blocked list?"),
    ("Buying process", "Who approves software purchases, and at what dollar thresholds?"),
    ("Buying process", "How long does a new vendor take to get through security, legal and procurement?"),
    ("Buying process", "Do you buy through AWS or Microsoft Marketplace today? Is a private marketplace enabled?"),
    ("Buying process", "Do marketplace purchases count toward your cloud commitment, and at what rate?"),
    ("Usage", "Where can we get license and active-user counts (admin consoles, SSO logs, a SaaS-management tool)?"),
    ("Usage", "Which tools have lots of licenses but few active users?"),
    ("Security & compliance", "Which tools hold regulated data (PII, PHI, PCI), and which frameworks apply (SOC 2, HIPAA, PCI, ISO 27001)?"),
    ("Next steps", "Who else should we talk to, and what would make this review a success for you?"),
]


def build(out: Path, example: bool) -> None:
    cat = yaml.safe_load(CATALOG.read_text(encoding="utf-8"))["entries"]
    wb = Workbook()

    # ---------------- Lists (hidden)
    ls = wb.active
    ls.title = "Lists"
    names = {}
    for i, (name, vals) in enumerate(LISTS.items(), 1):
        col = get_column_letter(i)
        ls[f"{col}1"] = name
        for j, v in enumerate(vals, 2):
            ls[f"{col}{j}"] = v
        names[name] = f"Lists!${col}$2:${col}${len(vals) + 1}"
    ls.sheet_state = "hidden"

    # ---------------- Lookup (hidden): every name and alias -> catalog product
    lk = wb.create_sheet("Lookup")
    lk["A1"], lk["B1"] = "Name or alias", "Catalog product"
    seen, j = set(), 2
    for e in cat:
        for nm in [e["product"], e.get("vendor", "") + " " + e["product"], *e.get("aliases", [])]:
            key = nm.strip().lower()
            if key and key not in seen:
                seen.add(key)
                lk[f"A{j}"], lk[f"B{j}"] = nm.strip(), e["product"]
                j += 1
    lk.sheet_state = "hidden"

    # ---------------- Instructions
    ins = wb.create_sheet("Instructions", 0)
    ins.sheet_view.showGridLines = False
    ins.column_dimensions["A"].width = 3
    ins.column_dimensions["B"].width = 26
    ins.column_dimensions["C"].width = 92
    ins["B2"] = "ISV Discovery Workbook"
    ins["B2"].font = f(20, True, TEAL)
    ins["B3"] = ("Document every software product your company pays for: what it's used for, when it renews, and "
                 "whether you'd consider alternatives. About 30–60 minutes for most companies.")
    ins["B3"].font = f(11, color="5A6A65")
    ins.merge_cells("B3:C3")
    ins["B3"].alignment = WRAP
    ins.row_dimensions[3].height = 32
    rows = [
        ("How to fill it in", None),
        ("1. Company Profile", "Your company, contacts, how you buy software, and any AWS or Microsoft cloud commitments."),
        ("2. ISV Inventory", "One row per product per contract. Pick the product from the dropdown, or type it if it isn't listed. "
                             "Fill the yellow columns; grey columns calculate themselves."),
        ("3. Discovery Questions", "Answer what you can. We'll cover the rest in the discovery session."),
        ("Then check", "Renewal Calendar lists every renewal in date order. Summary shows totals, renewals by month, "
                       "overlapping tools, and how open you are to alternatives."),
        ("Legend", None),
        ("Yellow cell", "You fill this in."),
        ("Grey cell", "Calculated or looked up from the catalog. Don't type here."),
        ("* in a header", "Required: product, primary use case, renewal date, and whether you're open to alternatives."),
        ("Red / amber highlight", "Renewal is past due or within 90 days; auto-renew notice deadline is within 45 days; "
                                  "fewer than 75% of licenses are in use."),
        ("Tips", None),
        ("Cost", "Give either the annual cost, or price per user per month plus number of licenses. Use the currency you pay in (assumed USD)."),
        ("Dates", "Type dates as YYYY-MM-DD (2027-01-31). If the renewal date is unknown, put your best estimate and say so in Notes."),
        ("Open to alternatives?", "Be candid. “Maybe – depends on price” is useful: it tells us to benchmark before the renewal."),
        ("Not in the list?", "Type the name anyway. The In catalog? column will say “No – will be reviewed”, and we'll add it."),
        ("Example row", "Row 6 of ISV Inventory is a filled-in example. Replace it or delete its contents."),
        ("What happens next", None),
        ("", "We review the workbook, match each product to our catalog of 157 ISVs, and send back a renewal plan: "
             "notice deadlines to calendar, licenses to right-size, overlapping tools, alternatives worth a look, "
             "and which renewals could be bought through AWS or Microsoft Marketplace against your cloud commitment."),
        ("Confidentiality", "This workbook contains commercial terms. Share it only with the people who need it."),
    ]
    r = 5
    for a, b in rows:
        if b is None:
            ins[f"B{r}"] = a
            ins[f"B{r}"].font = f(12, True, TEAL)
            r += 1
            continue
        ins[f"B{r}"] = a
        ins[f"B{r}"].font = f(10, True)
        ins[f"C{r}"] = b
        ins[f"C{r}"].font = f(10)
        ins[f"C{r}"].alignment = WRAP
        ins[f"B{r}"].alignment = WRAP
        if a == "Yellow cell":
            ins[f"B{r}"].fill = fill(INPUT)
        if a == "Grey cell":
            ins[f"B{r}"].fill = fill(CALC)
        ins.row_dimensions[r].height = 15 * max(1, len(b) // 95 + 1)
        r += 1
    ins[f"B{r + 1}"] = f"Catalog version: {len(cat)} products. Marketplace status as of October 2026; listings change, so treat it as a guide."
    ins[f"B{r + 1}"].font = f(9, italic=True, color="5A6A65")

    # ---------------- Company Profile
    pr = wb.create_sheet("Company Profile", 1)
    pr.sheet_view.showGridLines = False
    pr.column_dimensions["A"].width = 3
    pr.column_dimensions["B"].width = 46
    pr.column_dimensions["C"].width = 44
    pr["B2"] = "Company Profile"
    pr["B2"].font = f(16, True, TEAL)
    pr["B3"] = "Fill the yellow cells. * = required."
    pr["B3"].font = f(10, color="5A6A65")
    r = 5
    for label, key in PROFILE:
        if key is None:
            pr[f"B{r}"] = label
            pr[f"B{r}"].font = f(11, True, "FFFFFF")
            pr[f"B{r}"].fill = fill(TEAL)
            pr[f"C{r}"].fill = fill(TEAL)
            r += 1
            continue
        pr[f"B{r}"] = label
        pr[f"B{r}"].font = f(10)
        pr[f"B{r}"].border = BORDER
        cell = pr[f"C{r}"]
        cell.fill = fill(INPUT)
        cell.border = BORDER
        cell.font = f(10, color="0000FF")
        if key in PROFILE_FMT:
            cell.number_format = NUMFMT[PROFILE_FMT[key]]
        if key in PROFILE_LISTS:
            dv = DataValidation(type="list", formula1=f"={names[PROFILE_LISTS[key]]}", allow_blank=True)
            pr.add_data_validation(dv)
            dv.add(cell.coordinate)
        if example and key in PROFILE_EXAMPLE:
            cell.value = PROFILE_EXAMPLE[key]
        r += 1
    pr.freeze_panes = "B5"

    # ---------------- ISV Inventory
    inv = wb.create_sheet("ISV Inventory", 2)
    inv.sheet_view.zoomScale = 90
    inv["A1"] = "ISV Inventory"
    inv["A1"].font = f(16, True, TEAL)
    inv["A2"] = ("One row per product per contract. Yellow = you fill in · Grey = calculated · * = required. "
                 "Hover a header for help.")
    inv["A2"].font = f(10, color="5A6A65")
    inv["A3"] = "=\"Products listed: \"&COUNTA(B6:B255)&\"   ·   Annual spend: \"&TEXT(SUM(R6:R255),\"$#,##0\")"
    inv["A3"].font = f(10, True)
    # section bands
    start = 0
    for i, col in enumerate(COLS + [(None, None, "_end", 0, None, None, False, None)]):
        if i == len(COLS) or (i > 0 and col[2] != COLS[i - 1][2]):
            sec = COLS[start][2]
            a, b = get_column_letter(start + 1), get_column_letter(i)
            inv.merge_cells(f"{a}4:{b}4")
            inv[f"{a}4"] = sec
            inv[f"{a}4"].font = f(10, True, "FFFFFF")
            inv[f"{a}4"].fill = fill(BAND[sec])
            inv[f"{a}4"].alignment = CENTER
            start = i
    for i, (k, head, sec, w, kind, fmt, req, help_) in enumerate(COLS, 1):
        L = get_column_letter(i)
        c = inv[f"{L}{HEAD_ROW}"]
        c.value = head + (" *" if req else "")
        c.font = f(9, True, "FFFFFF")
        c.fill = fill(BAND[sec])
        c.alignment = CENTER
        c.border = BORDER
        if help_:
            c.comment = Comment(help_, "Workbook")
        inv.column_dimensions[L].width = w
    inv.row_dimensions[HEAD_ROW].height = 42
    inv.column_dimensions[C["key"]].hidden = True

    example_rows = FULL_EXAMPLE if example else [EXAMPLE_ROW]
    for r in range(FIRST, LAST + 1):
        data = example_rows[r - FIRST] if r - FIRST < len(example_rows) else {}
        for i, (k, head, sec, w, kind, fmt, req, help_) in enumerate(COLS, 1):
            cell = inv.cell(row=r, column=i)
            cell.border = BORDER
            cell.font = f(10, color="0000FF" if kind == "in" else INK)
            cell.alignment = Alignment(vertical="top", wrap_text=k in ("use_case", "pain", "notes", "alts", "barriers", "integrations", "other_uses", "considering"))
            if kind == "in":
                cell.fill = fill(INPUT)
                if k in data and data[k] is not None:
                    v = data[k]
                    cell.value = int(v) if k == "sat" else v
            else:
                cell.fill = fill(CALC)
                cell.value = formula(k, r)
            if fmt in NUMFMT:
                cell.number_format = NUMFMT[fmt]
        if r - FIRST < len(example_rows):
            inv.row_dimensions[r].height = 30
    inv.freeze_panes = f"C{FIRST}"
    inv.auto_filter.ref = f"A{HEAD_ROW}:{get_column_letter(len(COLS) - 1)}{LAST}"

    def dv_list(source, col, err=True, msg=None):
        dv = DataValidation(type="list", formula1=f"={source}", allow_blank=True, showErrorMessage=True)
        if not err:
            dv.errorStyle = "information"
            dv.error = msg or "Not in the list. That's fine; it will be reviewed."
            dv.errorTitle = "Not in the list"
        inv.add_data_validation(dv)
        dv.add(f"{C[col]}{FIRST}:{C[col]}{LAST}")

    dv_list(f"Catalog!$A$5:$A${4 + len(cat)}", "product", err=False,
            msg="This product isn't in our catalog yet. Keep it: the In catalog? column will flag it for review.")
    for k, lst in (("criticality", "Criticality"), ("sensitivity", "Sensitivity"), ("billing", "Billing"), ("channel", "Channel"),
                   ("auto", "YesNo"), ("tfc", "YesNo"), ("open", "Openness"), ("reason", "Reason"), ("outcome", "Outcome")):
        dv_list(names[lst], k)
    sat = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True, showErrorMessage=True,
                         error="Enter a whole number from 1 (very unhappy) to 5 (very satisfied).", errorTitle="Satisfaction")
    inv.add_data_validation(sat)
    sat.add(f"{C['sat']}{FIRST}:{C['sat']}{LAST}")
    for k in ("licenses", "active", "term", "notice"):
        dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True, showErrorMessage=True,
                            error="Enter a whole number.", errorTitle="Number")
        inv.add_data_validation(dv)
        dv.add(f"{C[k]}{FIRST}:{C[k]}{LAST}")
    for k in ("start", "renewal"):
        dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True, showErrorMessage=True,
                            error="Enter a date as YYYY-MM-DD.", errorTitle="Date")
        inv.add_data_validation(dv)
        dv.add(f"{C[k]}{FIRST}:{C[k]}{LAST}")

    rng = lambda k: f"{C[k]}{FIRST}:{C[k]}{LAST}"
    red, amber, green = fill("F7DCD9"), fill("FBE9D0"), fill(TEAL_SOFT)
    d0 = f"{C['days']}{FIRST}"
    inv.conditional_formatting.add(rng("days"), FormulaRule(formula=[f'AND(ISNUMBER({d0}),{d0}<0)'], fill=red))
    inv.conditional_formatting.add(rng("days"), FormulaRule(formula=[f'AND(ISNUMBER({d0}),{d0}<=90)'], fill=amber))
    inv.conditional_formatting.add(rng("window"), FormulaRule(formula=[f'{C["window"]}{FIRST}="Past due"'], fill=red))
    inv.conditional_formatting.add(rng("window"), FormulaRule(formula=[f'{C["window"]}{FIRST}="0–90 days"'], fill=amber))
    nb = f"{C['notice_by']}{FIRST}"
    inv.conditional_formatting.add(rng("notice_by"), FormulaRule(
        formula=[f'AND({C["auto"]}{FIRST}="Yes",ISNUMBER({nb}),{nb}-TODAY()<=45)'], fill=red))
    u0 = f"{C['util']}{FIRST}"
    inv.conditional_formatting.add(rng("util"), FormulaRule(formula=[f'AND(ISNUMBER({u0}),{u0}<0.75)'], fill=amber))
    o0 = f"{C['open']}{FIRST}"
    inv.conditional_formatting.add(rng("open"), FormulaRule(formula=[f'LEFT({o0},3)="Yes"'], fill=green, font=Font(name=FONT, color=TEAL, bold=True)))
    inv.conditional_formatting.add(rng("in_cat"), FormulaRule(formula=[f'LEFT({C["in_cat"]}{FIRST},2)="No"'], fill=amber))
    s0 = f"{C['sat']}{FIRST}"
    inv.conditional_formatting.add(rng("sat"), FormulaRule(formula=[f'AND(ISNUMBER({s0}),{s0}<=2)'], fill=amber))

    # ---------------- Renewal Calendar
    cal = wb.create_sheet("Renewal Calendar", 3)
    cal.sheet_view.showGridLines = False
    cal["A1"] = "Renewal Calendar"
    cal["A1"].font = f(16, True, TEAL)
    cal["A2"] = "Every product with a renewal date, soonest first. Calculated from ISV Inventory; don't type here."
    cal["A2"].font = f(10, color="5A6A65")
    heads = [("#", 5, None), ("Renewal date", 13, "date"), ("Days", 8, "int"), ("Window", 13, None), ("Product", 28, None),
             ("Business owner", 18, None), ("Annual spend ($)", 14, "money"), ("Auto-renews?", 11, None), ("Notice deadline", 13, "date"),
             ("Open to alternatives?", 24, None), ("Desired outcome", 24, None), ("AWS Marketplace", 13, None), ("Microsoft Marketplace", 13, None)]
    src = {"Renewal date": "renewal", "Days": "days", "Window": "window", "Product": "product", "Business owner": "owner",
           "Annual spend ($)": "annual_spend", "Auto-renews?": "auto", "Notice deadline": "notice_by",
           "Open to alternatives?": "open", "Desired outcome": "outcome", "AWS Marketplace": "aws", "Microsoft Marketplace": "ms"}
    for i, (h, w, fmt) in enumerate(heads, 1):
        L = get_column_letter(i)
        c = cal[f"{L}4"]
        c.value = h
        c.font = f(9, True, "FFFFFF")
        c.fill = fill(TEAL)
        c.alignment = CENTER
        cal.column_dimensions[L].width = w
    cal.row_dimensions[4].height = 30
    keyrng = f"'ISV Inventory'!${C['key']}${FIRST}:${C['key']}${LAST}"
    for n in range(1, NROWS + 1):
        r = 4 + n
        cal[f"A{r}"] = f'=IF(COUNT({keyrng})<{n},"",{n})'
        rowref = f"MATCH(SMALL({keyrng},$A{r}),{keyrng},0)"
        for i, (h, w, fmt) in enumerate(heads[1:], 2):
            L = get_column_letter(i)
            col = C[src[h]]
            cal[f"{L}{r}"] = f'=IF($A{r}="","",IFERROR(INDEX(\'ISV Inventory\'!${col}${FIRST}:${col}${LAST},{rowref})&"",""))' \
                if fmt is None else f'=IF($A{r}="","",IFERROR(INDEX(\'ISV Inventory\'!${col}${FIRST}:${col}${LAST},{rowref}),""))'
            if fmt:
                cal[f"{L}{r}"].number_format = NUMFMT[fmt]
        for i in range(1, len(heads) + 1):
            cal.cell(row=r, column=i).font = f(10)
            cal.cell(row=r, column=i).border = BORDER
    cal.freeze_panes = "A5"
    cal.conditional_formatting.add(f"C5:C{4 + NROWS}", FormulaRule(formula=['AND(ISNUMBER(C5),C5<0)'], fill=red))
    cal.conditional_formatting.add(f"C5:C{4 + NROWS}", FormulaRule(formula=['AND(ISNUMBER(C5),C5<=90)'], fill=amber))
    cal.conditional_formatting.add(f"J5:J{4 + NROWS}", FormulaRule(formula=['LEFT(J5,3)="Yes"'], fill=green))

    # ---------------- Summary
    sm = wb.create_sheet("Summary", 4)
    sm.sheet_view.showGridLines = False
    sm.column_dimensions["A"].width = 3
    for L, w in zip("BCDEFGH", (38, 14, 16, 4, 34, 14, 16)):
        sm.column_dimensions[L].width = w
    sm["B2"] = "Summary"
    sm["B2"].font = f(16, True, TEAL)
    sm["B3"] = "=IF('Company Profile'!C6=\"\",\"Calculated from ISV Inventory\",'Company Profile'!C6&\" · calculated from ISV Inventory\")"
    sm["B3"].font = f(10, color="5A6A65")
    I = lambda k: f"'ISV Inventory'!${C[k]}${FIRST}:${C[k]}${LAST}"

    def block(top, left, title, rows_, heads_=("Products", "Annual spend ($)")):
        L0, L1, L2 = left, chr(ord(left) + 1), chr(ord(left) + 2)
        sm[f"{L0}{top}"] = title
        sm[f"{L0}{top}"].font = f(11, True, "FFFFFF")
        for L in (L0, L1, L2):
            sm[f"{L}{top}"].fill = fill(TEAL)
        sm[f"{L1}{top}"] = heads_[0]
        sm[f"{L2}{top}"] = heads_[1]
        for L in (L1, L2):
            sm[f"{L}{top}"].font = f(9, True, "FFFFFF")
            sm[f"{L}{top}"].alignment = Alignment(horizontal="right")
        for j, (label, cnt, amt, fmt2) in enumerate(rows_, 1):
            r = top + j
            sm[f"{L0}{r}"] = label
            sm[f"{L1}{r}"] = cnt
            sm[f"{L2}{r}"] = amt
            sm[f"{L1}{r}"].number_format = "#,##0" if fmt2 != "pct" else "0%"
            sm[f"{L2}{r}"].number_format = NUMFMT["money"] if fmt2 != "pct2" else "0%"
            for L in (L0, L1, L2):
                sm[f"{L}{r}"].font = f(10)
                sm[f"{L}{r}"].border = BORDER
        return top + len(rows_) + 2

    P, R, D, S, O, A, OUT = I("product"), I("renewal"), I("days"), I("annual_spend"), I("open"), I("auto"), I("outcome")
    NB = I("notice_by")
    nxt = block(5, "B", "At a glance", [
        ("All products", f"=COUNTA({P})", f"=SUM({S})", None),
        ("Renewal past due", f'=COUNTIFS({D},"<0")', f'=SUMIFS({S},{D},"<0")', None),
        ("Renewing in 0–90 days", f'=COUNTIFS({D},">=0",{D},"<=90")', f'=SUMIFS({S},{D},">=0",{D},"<=90")', None),
        ("Renewing in 91–180 days", f'=COUNTIFS({D},">90",{D},"<=180")', f'=SUMIFS({S},{D},">90",{D},"<=180")', None),
        ("Renewing in 181–365 days", f'=COUNTIFS({D},">180",{D},"<=365")', f'=SUMIFS({S},{D},">180",{D},"<=365")', None),
        ("Auto-renew notice due within 45 days", f'=COUNTIFS({A},"Yes",{NB},">="&TODAY(),{NB},"<="&(TODAY()+45))',
         f'=SUMIFS({S},{A},"Yes",{NB},">="&TODAY(),{NB},"<="&(TODAY()+45))', None),
        ("Not in catalog (to review)", f'=COUNTIF({I("in_cat")},"No*")', f'=SUMIFS({S},{I("in_cat")},"No*")', None),
        ("Missing a renewal date", f'=SUMPRODUCT(({P}<>"")*({R}=""))', "", None),
    ])
    nxt2 = block(5, "F", "Open to alternatives after renewal?",
                 [(o, f'=COUNTIF({O},"{o}")', f'=SUMIFS({S},{O},"{o}")', None) for o in LISTS["Openness"]]
                 + [("Not answered", f'=SUMPRODUCT(({P}<>"")*({O}=""))', f'=SUMPRODUCT(({P}<>"")*({O}=""),{S})', None)])
    nxt3 = block(max(nxt, nxt2), "F", "Desired outcome at renewal",
                 [(o, f'=COUNTIF({OUT},"{o}")', f'=SUMIFS({S},{OUT},"{o}")', None) for o in LISTS["Outcome"]])
    # renewals by month, next 12 months
    mstart = 'DATE(YEAR(TODAY()),MONTH(TODAY()),1)'
    month_rows = []
    for k in range(12):
        a, b = f"EDATE({mstart},{k})", f"EDATE({mstart},{k + 1})"
        month_rows.append((f'=TEXT({a},"mmm yyyy")', f'=COUNTIFS({R},">="&{a},{R},"<"&{b})', f'=SUMIFS({S},{R},">="&{a},{R},"<"&{b})', None))
    nxt4 = block(nxt, "B", "Renewals by month (next 12 months)", month_rows)
    # overlapping categories
    cats = sorted({e["category"] for e in cat})
    CAT = I("cat_category")
    top = max(nxt3, nxt4)
    nxt5 = block(top, "B", "Tools per category (2+ = overlap)",
                 [(c_, f'=COUNTIF({CAT},"{c_}")', f'=SUMIFS({S},{CAT},"{c_}")', None) for c_ in cats])
    sm.conditional_formatting.add(f"C{top + 1}:C{top + len(cats)}", CellIsRule(operator="greaterThanOrEqual", formula=["2"], fill=amber, font=Font(name=FONT, bold=True)))
    sm.conditional_formatting.add(f"C{top + 1}:C{top + len(cats)}", CellIsRule(operator="equal", formula=["0"], font=Font(name=FONT, color="B0B8B5")))

    # ---------------- Discovery Questions
    dq = wb.create_sheet("Discovery Questions", 5)
    dq.sheet_view.showGridLines = False
    dq.column_dimensions["A"].width = 4
    dq.column_dimensions["B"].width = 20
    dq.column_dimensions["C"].width = 70
    dq.column_dimensions["D"].width = 70
    dq["B1"] = "Discovery Questions"
    dq["B1"].font = f(16, True, TEAL)
    dq["B2"] = "Answer what you can before the session. Short answers are fine."
    dq["B2"].font = f(10, color="5A6A65")
    for i, h in enumerate(("#", "Topic", "Question", "Your answer"), 1):
        c = dq.cell(row=4, column=i, value=h)
        c.font = f(10, True, "FFFFFF")
        c.fill = fill(TEAL)
    for n, (topic, q) in enumerate(QUESTIONS, 1):
        r = 4 + n
        dq[f"A{r}"] = n
        dq[f"B{r}"] = topic
        dq[f"C{r}"] = q
        dq[f"D{r}"].fill = fill(INPUT)
        dq[f"D{r}"].font = f(10, color="0000FF")
        for L in "ABCD":
            dq[f"{L}{r}"].border = BORDER
            dq[f"{L}{r}"].alignment = WRAP
            if L != "D":
                dq[f"{L}{r}"].font = f(10)
        dq.row_dimensions[r].height = 42
    dq.freeze_panes = "A5"

    # ---------------- Catalog
    cs = wb.create_sheet("Catalog", 6)
    cs["A1"] = "ISV Catalog"
    cs["A1"].font = f(16, True, TEAL)
    cs["A2"] = (f"{len(cat)} products. Marketplace status checked October 2026 (Listed = an ISV-published offer was found; "
                "Add-in only = a Teams/Office add-in or Entra app, not a purchasable offer; Unverified = not yet checked). "
                "Read-only reference used by the ISV Inventory dropdown and lookups.")
    cs["A2"].font = f(9, color="5A6A65")
    for i, (h, w) in enumerate((("Product", 32), ("Vendor", 20), ("Category", 24), ("AWS Marketplace", 15),
                                ("Microsoft Marketplace", 15), ("Alternatives", 70)), 1):
        L = get_column_letter(i)
        c = cs[f"{L}4"]
        c.value = h
        c.font = f(10, True, "FFFFFF")
        c.fill = fill(TEAL)
        cs.column_dimensions[L].width = w
    for j, e in enumerate(sorted(cat, key=lambda e: e["product"].lower()), 5):
        vals = (e["product"], e["vendor"], e["category"], STATUS.get(e["aws_marketplace"], e["aws_marketplace"]),
                STATUS.get(e["azure_marketplace"], e["azure_marketplace"]), ", ".join(e["alternatives"]))
        for i, v in enumerate(vals, 1):
            c = cs.cell(row=j, column=i, value=v)
            c.font = f(10)
    cs.freeze_panes = "A5"
    cs.auto_filter.ref = f"A4:F{4 + len(cat)}"
    cs.protection.sheet = True
    cs.protection.autoFilter = False
    cs.protection.sort = False

    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = {"Instructions": TEAL, "Company Profile": "2F5D7C", "ISV Inventory": "A35A00",
                                        "Renewal Calendar": "6B4E9B", "Summary": "9B2C48", "Discovery Questions": "2F5D7C",
                                        "Catalog": "5A6A65"}.get(ws.title, "FFFFFF")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--example", action="store_true", help="fill with a fictional customer")
    ap.add_argument("--out", type=Path)
    a = ap.parse_args()
    name = "ISV_Discovery_Workbook_EXAMPLE.xlsx" if a.example else "ISV_Discovery_Workbook.xlsx"
    build(a.out or ROOT / "templates" / name, a.example)
